from telegram import InlineKeyboardButton, InlineKeyboardMarkup
from decouple import config
from dotenv import load_dotenv
from telegram import Bot, Update
from telegram.ext import CommandHandler, CallbackContext, Updater
from mutagen.easyid3 import EasyID3
from mutagen.mp3 import MP3
from mutagen.mp4 import MP4
from mutagen.flac import FLAC
from moviepy.editor import VideoFileClip
from PIL import Image
import docx
import os
from subliminal import download_best_subtitles, region
from pathlib import Path
import json
import openpyxl
import csv
from PyPDF2 import PdfReader
from pptx import Presentation
from zipfile import ZipFile
import sqlite3
import xml.etree.ElementTree as ET
import markdown

# Load environment variables from a .env file
load_dotenv()

# Use decouple to get the bot token from the environment
bot_token = config('BOT_TOKEN')

bot = Bot(token=bot_token)

def start(update, context):
    user = update.effective_user
    context.bot.send_photo(
        chat_id=update.effective_chat.id,
        photo="https://example.com/your_photo.jpg",  # Replace with the actual photo URL
        caption=f"Hello {user.first_name}!",
        reply_markup=get_start_inline_keyboard(),
    )

def get_start_inline_keyboard():
    keyboard = [
        [
            InlineKeyboardButton("Owner", url='https://t.me/stupidboi69'),
            InlineKeyboardButton("Channel", url='https://t.me/anime_downlord'),
        ],
        [InlineKeyboardButton("Group", url='https://t.me/anime_download_group')],
    ]
    return InlineKeyboardMarkup(keyboard)

def get_audio_metadata(file_path):
    try:
        if file_path.lower().endswith(".mp3"):
            audio = MP3(file_path)
        elif file_path.lower().endswith((".mp4", ".m4a")):
            audio = MP4(file_path)
        elif file_path.lower().endswith(".flac"):
            audio = FLAC(file_path)
        else:
            return "Unsupported audio format"

        title = audio.get("title", [""])[0]
        artist = audio.get("artist", [""])[0]
        album = audio.get("album", [""])[0]
        year = audio.get("date", [""])[0]
        genre = audio.get("genre", [""])[0]
        track_number = audio.get("tracknumber", [""])[0]
        comment = audio.get("comment", [""])[0]
        bitrate = audio.info.bitrate
        duration_seconds = audio.info.length

        return {
            "Title": title,
            "Artist": artist,
            "Album": album,
            "Year": year,
            "Genre": genre,
            "Track Number": track_number,
            "Comment": comment,
            "Bitrate": bitrate,
            "Duration (seconds)": duration_seconds
        }
    except Exception as e:
        return f"Error: {e}"

def get_video_metadata(file_path):
    try:
        video = VideoFileClip(file_path)
        duration = video.duration
        resolution = video.size
        fps = video.fps
        return {
            "Duration": duration,
            "Resolution": resolution,
            "FPS": fps
        }
    except Exception as e:
        return f"Error: {e}"

def get_document_metadata(file_path):
    try:
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs]
        return {
            "Paragraphs": paragraphs
        }
    except Exception as e:
        return f"Error: {e}"

def get_subtitle_metadata(file_path):
    try:
        subtitles = download_best_subtitles([file_path], {region.Region.subtitles, 'en'})
        return {
            "Subtitles": [subtitle.to_dict() for subtitle in subtitles]
        }
    except Exception as e:
        return f"Error: {e}"

def get_image_metadata(file_path):
    try:
        image = Image.open(file_path)
        width, height = image.size
        format_type = image.format
        return {
            "Width": width,
            "Height": height,
            "Format": format_type
        }
    except Exception as e:
        return f"Error: {e}"

def get_pdf_metadata(file_path):
    try:
        with open(file_path, 'rb') as file:
            pdf_reader = PdfReader(file)
            num_pages = len(pdf_reader.pages)
            return {
                "Number of Pages": num_pages
            }
    except Exception as e:
        return f"Error: {e}"

def get_excel_metadata(file_path):
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet_names = workbook.sheetnames
        return {
            "Sheet Names": sheet_names
        }
    except Exception as e:
        return f"Error: {e}"

def get_text_metadata(file_path):
    try:
        with open(file_path, 'r') as file:
            lines = file.readlines()
            num_lines = len(lines)
            return {
                "Number of Lines": num_lines
            }
    except Exception as e:
        return f"Error: {e}"

def get_csv_metadata(file_path):
    try:
        with open(file_path, 'r') as file:
            csv_reader = csv.reader(file)
            header = next(csv_reader)  # Assuming the first row is the header
            num_rows = sum(1 for row in csv_reader)
            return {
                "Header": header,
                "Number of Rows": num_rows
            }
    except Exception as e:
        return f"Error: {e}"

def get_json_metadata(file_path):
    try:
        with open(file_path, 'r') as file:
            data = json.load(file)
            return {
                "Data": data
            }
    except Exception as e:
        return f"Error: {e}"

def get_pptx_metadata(file_path):
    try:
        presentation = Presentation(file_path)
        num_slides = len(presentation.slides)
        return {
            "Number of Slides": num_slides
        }
    except Exception as e:
        return f"Error: {e}"

def get_zip_metadata(file_path):
    try:
        with ZipFile(file_path, 'r') as zip_file:
            file_list = zip_file.namelist()
            return {
                "Files in ZIP": file_list
            }
    except Exception as e:
        return f"Error: {e}"

def get_sqlite_metadata(file_path):
    try:
        connection = sqlite3.connect(file_path)
        cursor = connection.cursor()
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
        tables = cursor.fetchall()
        connection.close()
        return {
            "Tables in SQLite Database": tables
        }
    except Exception as e:
        return f"Error: {e}"

def get_xml_metadata(file_path):
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        elements = [elem.tag for elem in root.iter()]
        return {
            "XML Elements": elements
        }
    except Exception as e:
        return f"Error: {e}"

def get_md_metadata(file_path):
    try:
        with open(file_path, 'r') as file:
            content = file.read()
            html_content = markdown.markdown(content)
            return {
                "HTML Content": html_content
            }
    except Exception as e:
        return f"Error: {e}"

# Modify get_file_metadata function to include new file types
def get_file_metadata(file_path):
    try:
        file_name, file_extension = os.path.splitext(os.path.basename(file_path))
        if file_extension.lower() in (".mp3", ".mp4", ".m4a", ".flac"):
            return get_audio_metadata(file_path)
        elif file_extension.lower() in (".avi", ".mp4", ".mkv"):
            return get_video_metadata(file_path)
        elif file_extension.lower() == ".docx":
            return get_document_metadata(file_path)
        elif file_extension.lower() == ".srt":
            return get_subtitle_metadata(file_path)
        elif file_extension.lower() in (".jpg", ".jpeg", ".png", ".bmp", ".gif"):
            return get_image_metadata(file_path)
        elif file_extension.lower() == ".pdf":
            return get_pdf_metadata(file_path)
        elif file_extension.lower() in (".xlsx", ".xls"):
            return get_excel_metadata(file_path)
        elif file_extension.lower() == ".txt":
            return get_text_metadata(file_path)
        elif file_extension.lower() == ".csv":
            return get_csv_metadata(file_path)
        elif file_extension.lower() == ".json":
            return get_json_metadata(file_path)
        elif file_extension.lower() == ".pptx":
            return get_pptx_metadata(file_path)
        elif file_extension.lower() == ".zip":
            return get_zip_metadata(file_path)
        elif file_extension.lower() == ".sqlite":
            return get_sqlite_metadata(file_path)
        elif file_extension.lower() == ".xml":
            return get_xml_metadata(file_path)
        elif file_extension.lower() == ".md":
            return get_md_metadata(file_path)
        else:
            return "Unsupported file format"
    except Exception as e:
        return f"Error: {e}"

# ... (existing code)

if __name__ == "__main__":
    updater = Updater(token=bot_token, use_context=True)
    dispatcher = updater.dispatcher

    start_handler = CommandHandler('start', start)
    extract_metadata_handler = CommandHandler('extract_metadata', extract_metadata)

    dispatcher.add_handler(start_handler)
    dispatcher.add_handler(extract_metadata_handler)

    updater.start_polling()
    updater.idle()
